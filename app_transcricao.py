import os
import threading
import json
import tkinter.filedialog as fd
from tkinterdnd2 import DND_FILES, TkinterDnD
import customtkinter as ctk

try:
    import whisper
except ImportError:
    import tkinter.messagebox as mb
    mb.showerror("Erro", "Whisper não instalado. Rode o instalar_dependencias.bat!")
    exit()

ctk.set_appearance_mode("System")  # "Dark", "Light" ou "System"
ctk.set_default_color_theme("blue")

# --- Bloco 1: HISTÓRICO ---
HISTORICO_ARQUIVO = "historico_transcricoes.json"
MAX_HISTORICO = 10  # Quantos arquivos manter no histórico

def carregar_historico():
    if os.path.exists(HISTORICO_ARQUIVO):
        try:
            with open(HISTORICO_ARQUIVO, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return []
    return []

def salvar_historico(lista):
    with open(HISTORICO_ARQUIVO, "w", encoding="utf-8") as f:
        json.dump(lista, f, ensure_ascii=False, indent=2)

class TranscritorApp(ctk.CTkFrame):
    def __init__(self, master, **kwargs):
        super().__init__(master, **kwargs)
        self.master = master
        self.master.drop_target_register(DND_FILES)
        self.master.dnd_bind('<<Drop>>', self.drop_event)

        # Variáveis
        self.arquivo_escolhido = ""
        self.transcricao = ""

        # Interface
        self.frame = ctk.CTkFrame(self, fg_color="#f3f6fa", corner_radius=16)
        self.frame.pack(padx=24, pady=18, fill="both", expand=True)

        self.label_drop = ctk.CTkLabel(self.frame, text="Arraste e solte o arquivo aqui", font=("Arial", 12), text_color="#666")
        self.label_drop.pack(pady=(6,2))

        # TEMA
        temas = ["System", "Light", "Dark"]
        self.opcao_tema = ctk.CTkOptionMenu(
            master=self.frame,
            values=temas,
            command=self.trocar_tema,
            width=120
        )
        self.opcao_tema.set("System")
        self.opcao_tema.place(relx=1.0, rely=0.0, anchor="ne", x=-10, y=10)

        #cabeçalho titulo
        self.label = ctk.CTkLabel(
        self.frame,
        text="TRANSCRITOR UNIVERSAL",
        font=("Arial Black", 25, "bold"),
        text_color="#15386b")  # Azul bem escuro
        self.label.pack(pady=(18,10))

        self.btn_abrir = ctk.CTkButton(self.frame, text="Escolher arquivo", command=self.abrir_arquivo, width=220, height=36)
        self.btn_abrir.pack(pady=(8,8))

        self.lbl_arq = ctk.CTkLabel(self.frame, text="Nenhum arquivo selecionado.", text_color="#999")
        self.lbl_arq.pack(pady=(0,14))

        idiomas = [
            "auto",       # automático
            "pt",         # Português
            "en",         # Inglês
            "es",         # Espanhol
            "fr",         # Francês
            "de",         # Alemão
            "it",         # Italiano
            "ru",         # Russo
            "zh",         # Chinês
        ]
        self.idioma_var = ctk.StringVar(value="auto")
        self.opcao_idioma = ctk.CTkOptionMenu(
            self.frame,
            values=idiomas,
            variable=self.idioma_var,
            width=120
        )
        self.opcao_idioma.place(relx=0.0, rely=0.0, anchor="nw", x=10, y=10)

        self.btn_transcrever = ctk.CTkButton(self.frame, text="Iniciar Transcrição", command=self.iniciar_transcricao, width=220, height=38, state="disabled")
        self.btn_transcrever.pack(pady=8)

        self.txt_result = ctk.CTkTextbox(self.frame, height=120, font=("Consolas", 13), wrap="word")
        self.txt_result.pack(padx=12, pady=8, fill="both", expand=False)
        self.txt_result.insert("1.0", "Sua transcrição aparecerá aqui.")
        self.txt_result.configure(state="disabled")

        self.export_frame = ctk.CTkFrame(self.frame, fg_color="transparent")
        self.export_frame.pack(pady=10)

        self.btn_export_txt = ctk.CTkButton(self.export_frame, text="Exportar para TXT", command=lambda: self.exportar("txt"), state="disabled", width=110)
        self.btn_export_txt.grid(row=0, column=0, padx=8)

        self.btn_export_json = ctk.CTkButton(self.export_frame, text="Exportar para JSON", command=lambda: self.exportar("json"), state="disabled", width=110)
        self.btn_export_json.grid(row=0, column=1, padx=8)

        self.btn_export_md = ctk.CTkButton(self.export_frame, text="Exportar para Markdown", command=lambda: self.exportar("md"), state="disabled", width=110)
        self.btn_export_md.grid(row=0, column=2, padx=8)

        # Barra de Progresso (Indicação Visual)
        self.progress = ctk.CTkProgressBar(self.frame)
        self.progress.set(0)  # Começa zerada
        self.progress.pack(pady=(8,2))
        self.progress.pack_forget()  # Esconde no início

        self.lbl_status = ctk.CTkLabel(self.frame, text="", font=("Arial", 11), text_color="#1a73e8")
        self.lbl_status.pack(pady=6)

        # Atalhos de teclado
        self.master.bind_all("<Control-o>", lambda event: self.abrir_arquivo())
        self.master.bind_all("<Control-t>", lambda event: self.iniciar_transcricao())
        self.master.bind_all("<Control-e>", lambda event: self.exportar_via_atalho())
        self.master.bind_all("<Control-q>", lambda event: self.master.destroy())


        # Histórico de transcrições
        self.historico = carregar_historico()
        self.label_hist = ctk.CTkLabel(self.frame, text="Últimas transcrições:", font=("Arial", 13, "bold"), text_color="#222")
        self.label_hist.pack(pady=(10,0))

        self.lista_hist = ctk.CTkTextbox(self.frame, height=64, font=("Consolas", 11), wrap="word")
        self.lista_hist.pack(padx=10, pady=(4,8), fill="x")
        self.lista_hist.configure(state="disabled")
        self.atualizar_historico()

    def trocar_tema(self, value):
        ctk.set_appearance_mode(value)
        self.lbl_status.configure(text=f"Tema alterado para {value}")

    def atualizar_historico(self):
        self.lista_hist.configure(state="normal")
        self.lista_hist.delete("1.0", "end")
        if not self.historico:
            self.lista_hist.insert("1.0", "Nenhuma transcrição recente.")
        else:
            for item in reversed(self.historico[-MAX_HISTORICO:]):
                self.lista_hist.insert("end", f"- {item['arquivo']} ({item['tipo']})\n")
        self.lista_hist.configure(state="disabled")

    def abrir_arquivo(self):
        tipos = [
                ("Todos suportados", "*.mp3 *.wav *.m4a *.flac *.mp4 *.avi *.mov *.mkv *.txt *.pdf *.docx *.xlsx *.jpg *.jpeg *.png"),
                ("Áudio", "*.mp3 *.wav *.m4a *.flac"),
                ("Vídeo", "*.mp4 *.avi *.mov *.mkv"),
                ("Texto", "*.txt"),
                ("PDF", "*.pdf"),
                ("Word DOCX", "*.docx"),
                ("Planilha Excel", "*.xlsx"),
                ("Imagem", "*.jpg *.jpeg *.png"),
                ("Todos os arquivos", "*.*")
            ]
        arq = fd.askopenfilename(title="Escolha o arquivo para transcrever", filetypes=tipos)
        if arq:
            self.arquivo_escolhido = arq
            self.lbl_arq.configure(text=f"Selecionado: {os.path.basename(arq)}", text_color="#222")
            self.btn_transcrever.configure(state="normal")
            self.label_drop.pack_forget()
        else:
            self.lbl_arq.configure(text="Nenhum arquivo selecionado.", text_color="#999")
            self.btn_transcrever.configure(state="disabled")

    def iniciar_transcricao(self):
        self.btn_transcrever.configure(state="disabled")
        self.btn_abrir.configure(state="disabled")
        self.txt_result.configure(state="normal")
        self.txt_result.delete("1.0", "end")
        self.txt_result.insert("1.0", "Transcrevendo, aguarde...")
        self.txt_result.configure(state="disabled")
        self.lbl_status.configure(text="Processando...")
        self.progress.pack(pady=(8,2))
        self.progress.set(0.1)
        self.lbl_status.configure(text="Transcrevendo... Aguarde.")
        threading.Thread(target=self.transcrever_thread).start()

    def transcrever_thread(self):
        arq = self.arquivo_escolhido
        ext = os.path.splitext(arq)[1].lower()
        try:
            texto = ""
            if ext == ".txt":
                with open(arq, "r", encoding="utf-8") as f:
                    texto = f.read()
            elif ext in [".mp3", ".wav", ".m4a", ".flac", ".mp4", ".avi", ".mov", ".mkv"]:
                idioma = self.idioma_var.get()
                model = whisper.load_model("base")
                if idioma == "auto":
                    result = model.transcribe(arq)
                else:
                    result = model.transcribe(arq, language=idioma)
                texto = result['text']
            elif ext == ".pdf":
                import pdfplumber
                with pdfplumber.open(arq) as pdf:
                    texto = "\n".join(page.extract_text() or "" for page in pdf.pages)
            elif ext == ".docx":
                from docx import Document
                doc = Document(arq)
                texto = "\n".join([p.text for p in doc.paragraphs])
            elif ext == ".xlsx":
                import openpyxl
                wb = openpyxl.load_workbook(arq)
                texto = ""
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        texto += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
            elif ext in [".jpg", ".jpeg", ".png"]:
                from PIL import Image
                import pytesseract
                img = Image.open(arq)
                idioma_ocr = self.idioma_var.get()
                if idioma_ocr == "auto":
                    idioma_ocr = "por"
                texto = pytesseract.image_to_string(img, lang=idioma_ocr)
            else:
                texto = "Tipo de arquivo não suportado para transcrição."
            self.transcricao = texto.strip()
            self.progress.set(1)
            self.after(500, self.progress.pack_forget)
            self.exibir_resultado()
        except Exception as e:
            self.txt_result.configure(state="normal")
            self.txt_result.delete("1.0", "end")
            self.txt_result.insert("1.0", f"Erro: {str(e)}")
            self.txt_result.configure(state="disabled")
            self.lbl_status.configure(text="Erro ao transcrever.")
            self.progress.set(0)
            self.after(500, self.progress.pack_forget)
        self.btn_transcrever.configure(state="normal")
        self.btn_abrir.configure(state="normal")

    def exibir_resultado(self):
        self.txt_result.configure(state="normal")
        self.txt_result.delete("1.0", "end")
        self.txt_result.insert("1.0", self.transcricao)
        self.txt_result.configure(state="disabled")
        self.lbl_status.configure(text="Transcrição concluída com sucesso!")
        self.btn_export_txt.configure(state="normal")
        self.btn_export_json.configure(state="normal")
        self.btn_export_md.configure(state="normal")
    
        # Atualiza histórico ao concluir transcrição
        tipo = "texto"
        ext = os.path.splitext(self.arquivo_escolhido)[1].lower()
        if ext in [".mp3", ".wav", ".m4a", ".flac"]:
            tipo = "áudio"
        elif ext in [".mp4", ".avi", ".mov", ".mkv"]:
            tipo = "vídeo"
        elif ext in [".txt"]:
            tipo = "texto"
        self.historico.append({
            "arquivo": os.path.basename(self.arquivo_escolhido),
            "tipo": tipo
        })
        if len(self.historico) > MAX_HISTORICO:
            self.historico = self.historico[-MAX_HISTORICO:]
        salvar_historico(self.historico)
        self.atualizar_historico()

    def exportar(self, formato):
        if not self.transcricao.strip():
            return
        if formato == "txt":
            ext = ".txt"
            tipos = [("Arquivo TXT", "*.txt")]
            conteudo = self.transcricao
        elif formato == "json":
            ext = ".json"
            tipos = [("Arquivo JSON", "*.json")]
            conteudo = json.dumps({"transcricao": self.transcricao}, ensure_ascii=False, indent=2)
        elif formato == "md":
            ext = ".md"
            tipos = [("Markdown", "*.md")]
            conteudo = f"# Transcrição\n\n{self.transcricao}"
        else:
            return

        saida = fd.asksaveasfilename(defaultextension=ext, filetypes=tipos)
        if saida:
            with open(saida, "w", encoding="utf-8") as f:
                f.write(conteudo)
            self.lbl_status.configure(text=f"Arquivo salvo: {os.path.basename(saida)}")

    def exportar_via_atalho(self):
        import tkinter.simpledialog as sd
        opcoes = {"1": "txt", "2": "json", "3": "md"}
        escolha = sd.askstring("Exportar", "Escolha o formato:\n1 - TXT\n2 - JSON\n3 - Markdown", parent=self)
        if escolha and escolha.strip() in opcoes:
            self.exportar(opcoes[escolha.strip()])
    
    def drop_event(self, event):
        arq_drop = event.data.strip().replace("{", "").replace("}", "")
        if os.path.exists(arq_drop):
            self.arquivo_escolhido = arq_drop
            self.lbl_arq.configure(text=f"Selecionado (Drag & Drop): {os.path.basename(arq_drop)}", text_color="#222")
            self.btn_transcrever.configure(state="normal")
            self.label_drop.pack_forget()
        else:
            self.lbl_status.configure(text="Arquivo não encontrado para Drag & Drop.")

if __name__ == "__main__":
    root = TkinterDnD.Tk()
    root.title("Transcritor Universal Áudio/Vídeo/Text")
    root.geometry("800x650")
    root.resizable(True, True)
    app = TranscritorApp(root)
    app.pack(fill="both", expand=True)
    root.mainloop()
