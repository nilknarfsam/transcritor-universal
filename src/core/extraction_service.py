"""Extração de texto de documentos e OCR de imagens (sem Whisper)."""

from __future__ import annotations

from src.models.transcription_job import (
    DOCUMENT_EXTENSIONS,
    IMAGE_EXTENSIONS,
    TEXT_EXTENSIONS,
    TranscriptionJob,
)


class ExtractionService:
    """Extrai texto de TXT, PDF, DOCX, XLSX e imagens via Tesseract."""

    def extract(self, job: TranscriptionJob, language: str = "auto") -> str:
        """Retorna o texto extraído conforme a extensão do arquivo."""
        ext = job.extension
        path = job.file_path

        if ext in TEXT_EXTENSIONS:
            return self._extract_txt(path)
        if ext == ".pdf":
            return self._extract_pdf(path)
        if ext == ".docx":
            return self._extract_docx(path)
        if ext == ".xlsx":
            return self._extract_xlsx(path)
        if ext in IMAGE_EXTENSIONS:
            return self._extract_image(path, language)
        raise ValueError(f"Tipo de arquivo não suportado: {ext}")

    def can_extract(self, job: TranscriptionJob) -> bool:
        """Indica se o job pode ser processado por este serviço (não áudio/vídeo)."""
        return job.extension in (TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS)

    @staticmethod
    def _extract_txt(path: str) -> str:
        try:
            with open(path, encoding="utf-8") as f:
                return f.read().strip()
        except UnicodeDecodeError:
            with open(path, encoding="latin-1") as f:
                return f.read().strip()

    @staticmethod
    def _extract_pdf(path: str) -> str:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n".join(pages).strip()

    @staticmethod
    def _extract_docx(path: str) -> str:
        from docx import Document

        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs).strip()

    @staticmethod
    def _extract_xlsx(path: str) -> str:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    line = "\t".join(
                        str(cell) if cell is not None else "" for cell in row
                    )
                    lines.append(line)
            return "\n".join(lines).strip()
        finally:
            wb.close()

    @staticmethod
    def _extract_image(path: str, language: str) -> str:
        from PIL import Image
        import pytesseract

        ocr_lang = language if language != "auto" else "por"
        with Image.open(path) as img:
            return pytesseract.image_to_string(img, lang=ocr_lang).strip()
